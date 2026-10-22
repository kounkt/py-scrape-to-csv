import csv

from openpyxl import load_workbook

from scraper.exporter import write_csv, write_xlsx

FIELDS = ["title", "price"]
ROWS = [{"title": "本の題名", "price": 51.77}, {"title": "=HYPERLINK(\"x\")", "price": None, "extra": "無視"}]


def test_write_csv_bom_and_formula_guard(tmp_path):
    out = tmp_path / "sub" / "a.csv"
    assert write_csv(ROWS, out, FIELDS) == 2
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # Excel 用の BOM
    with out.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0] == {"title": "本の題名", "price": "51.77"}
    assert rows[1]["title"].startswith("'=")  # 数式として動かない
    assert "extra" not in rows[1]


def test_write_xlsx(tmp_path):
    out = tmp_path / "a.xlsx"
    assert write_xlsx(ROWS, out, FIELDS) == 2
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == FIELDS
    assert ws["A2"].value == "本の題名"
    assert ws["B2"].value == 51.77
    assert ws["A3"].value == '=HYPERLINK("x")'  # ' は付けず、型を文字列に固定する
    assert ws["A3"].data_type == "s"
    assert ws.freeze_panes == "A2"


def test_write_xlsx_strips_control_characters(tmp_path):
    out = tmp_path / "b.xlsx"
    write_xlsx([{"title": "a\x00b\x1fc", "price": 1}], out, FIELDS)
    assert load_workbook(out).active["A2"].value == "abc"
