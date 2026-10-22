import csv

from openpyxl import load_workbook

from scraper.cli import main
from scraper.fetcher import PoliteFetcher
from tests.conftest import FakeResponse, FakeSession

SITE = "https://books.example.test"


def site_routes(fixture_html):
    return {
        SITE + "/robots.txt": FakeResponse(200, "User-agent: *\nDisallow: /admin/\n"),
        SITE + "/": FakeResponse(200, fixture_html("list_page1.html")),
        SITE + "/catalogue/page-2.html": FakeResponse(200, fixture_html("list_page2.html")),
    }


def fetcher_for(routes, clock):
    return PoliteFetcher("test-bot/1.0", min_interval=1.0, session=FakeSession(routes), sleep=clock.sleep, clock=clock.clock)


def test_end_to_end_csv(tmp_path, fixture_html, fake_clock):
    out = tmp_path / "books.csv"
    code = main(["--start-url", SITE + "/", "--max-pages", "5", "--out", str(out)], fetcher=fetcher_for(site_routes(fixture_html), fake_clock))
    assert code == 0
    with out.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["title"] for r in rows] == ["Sample Book One", "'=SUM(A1:A2)", "Sample Book Three"]
    assert rows[0]["in_stock"] == "True"


def test_max_pages_limits_crawl(tmp_path, fixture_html, fake_clock):
    out = tmp_path / "books.xlsx"
    code = main(["--start-url", SITE + "/", "--max-pages", "1", "--out", str(out)], fetcher=fetcher_for(site_routes(fixture_html), fake_clock))
    assert code == 0
    assert load_workbook(out).active.max_row == 3  # 見出し＋2件


def test_disallowed_start_url_returns_3(tmp_path, fixture_html, fake_clock):
    code = main(["--start-url", SITE + "/admin/", "--out", str(tmp_path / "x.csv")], fetcher=fetcher_for(site_routes(fixture_html), fake_clock))
    assert code == 3


def test_bad_extension_returns_2(tmp_path):
    assert main(["--start-url", SITE + "/", "--out", str(tmp_path / "x.json")]) == 2


def test_zero_items_returns_5(tmp_path, fake_clock):
    routes = {SITE + "/": FakeResponse(200, "<html><body>no items</body></html>")}
    out = tmp_path / "x.csv"
    assert main(["--start-url", SITE + "/", "--out", str(out)], fetcher=fetcher_for(routes, fake_clock)) == 5
    assert not out.exists()


def test_next_page_on_other_host_is_not_followed(tmp_path, fixture_html, fake_clock):
    html = fixture_html("list_page1.html").replace('href="catalogue/page-2.html"', 'href="https://other.example.test/page-2.html"')
    routes = {SITE + "/": FakeResponse(200, html)}
    fetcher = fetcher_for(routes, fake_clock)
    out = tmp_path / "x.csv"
    assert main(["--start-url", SITE + "/", "--max-pages", "3", "--out", str(out)], fetcher=fetcher) == 0
    assert all("other.example.test" not in url for url in fetcher.session.calls)


def test_bad_arguments_return_2(tmp_path):
    assert main(["--start-url", SITE + "/", "--delay", "-1", "--out", str(tmp_path / "x.csv")]) == 2
    assert main(["--start-url", "ftp://example.test/", "--out", str(tmp_path / "x.csv")]) == 2
