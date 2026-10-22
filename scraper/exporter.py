"""CSV / Excel への書き出し。"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(value):
    """Excel で開いた時に数式として動かないよう、記号で始まる文字列の先頭に ' を付ける。"""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _xlsx_cell(value):
    """Excel 用：制御文字を取り除く（数式にしない処理は書き込み後に型を「文字列」に固定して行う）。"""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def write_csv(rows: Iterable[dict], path: str | Path, fieldnames: Sequence[str]) -> int:
    """UTF-8（BOM 付き）で書く。Excel でそのまま開いても文字化けしない。書いた行数を返す。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: _safe_cell(row.get(k)) for k in fieldnames})
            count += 1
    return count


def write_xlsx(rows: Iterable[dict], path: str | Path, fieldnames: Sequence[str], sheet_title: str = "data") -> int:
    """Excel 形式で書く。見出しの固定と列幅の調整つき。書いた行数を返す。

    文字列のセルは型を「文字列」に固定するので、= で始まっても数式として動かない（' は付けない）。
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(fieldnames))
    count = 0
    widths = [len(name) for name in fieldnames]
    for row in rows:
        values = [_xlsx_cell(row.get(k)) for k in fieldnames]
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        for i, v in enumerate(values):
            widths[i] = max(widths[i], min(len(str(v)) if v is not None else 0, 60))
        count += 1
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w + 2
    ws.freeze_panes = "A2"
    wb.save(path)
    return count
